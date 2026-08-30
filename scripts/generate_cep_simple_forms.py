from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_DIR = Path("docs/entregables/cep-recogida-datos-simplificados")
CENTERS = (
    ("CEP Norte", "cep-norte"),
    ("CEP Santa Cruz", "cep-santa-cruz"),
    ("CEP Sur", "cep-sur"),
)

BLUE = "17365D"
LIGHT_BLUE = "EAF1F8"
LIGHT_GRAY = "F3F4F6"
BORDER = "B9C3CF"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(90, 100, 112)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm):
    width_dxa = int(width_cm / 2.54 * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_cm):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_dxa = int(sum(widths_cm) / 2.54 * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_cm in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_cm[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=10, bold=False, color=TEXT):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def style_paragraph(paragraph, after=0, before=0, line=1.0, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def configure_document(landscape=False):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(29.7 if landscape else 21.0)
    section.page_height = Cm(21.0 if landscape else 29.7)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    return document


def add_title(document, title, subtitle):
    paragraph = document.add_paragraph()
    style_paragraph(paragraph, after=2)
    run = paragraph.add_run(title)
    set_run(run, size=18, bold=True, color=RGBColor.from_string(BLUE))

    paragraph = document.add_paragraph()
    style_paragraph(paragraph, after=9)
    run = paragraph.add_run(subtitle)
    set_run(run, size=9.5, color=MUTED)


def style_table_text(table, header_rows=0, label_columns=()):
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                style_paragraph(paragraph, line=1.0)
                for run in paragraph.runs:
                    is_header = row_index < header_rows
                    is_label = column_index in label_columns
                    set_run(run, size=8.5 if is_header else 9, bold=is_header or is_label)
                    if is_header:
                        run.font.color.rgb = RGBColor(255, 255, 255)


def build_company_form(center_name, output_path):
    document = configure_document(landscape=False)
    add_title(document, f"Datos de empresa · {center_name}", "Rellene los datos actuales de la empresa responsable del centro.")

    fields = (
        "Razón social completa",
        "Nombre comercial",
        "NIF / CIF",
        "Forma jurídica y situación legal",
        "Domicilio fiscal",
        "Dirección física del centro",
        "Código postal, municipio y provincia",
        "Teléfono principal",
        "Correo electrónico general",
        "Correo electrónico de facturación",
        "Página web",
        "Persona de contacto",
        "Cargo",
        "Teléfono y correo de contacto",
    )
    table = document.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, (5.2, 12.5))

    for index, label in enumerate(fields):
        label_cell, value_cell = table.rows[index].cells
        set_cell_shading(label_cell, LIGHT_BLUE)
        label_cell.text = label
        value_cell.text = " "
        label_cell.paragraphs[0].paragraph_format.keep_with_next = False
        value_cell.paragraphs[0].paragraph_format.keep_with_next = False
    style_table_text(table, label_columns=(0,))

    note = document.add_paragraph()
    style_paragraph(note, before=7)
    run = note.add_run("Programa de contabilidad utilizado (si procede): ")
    set_run(run, size=9, bold=True)
    run = note.add_run("____________________________________________________________")
    set_run(run, size=9, color=MUTED)
    document.save(output_path)


def build_staff_form(center_name, output_path):
    document = configure_document(landscape=True)
    add_title(document, f"Datos de personal y accesos · {center_name}", "Incluya únicamente a las personas que necesitarán acceso al sistema.")

    headers = ("Departamento", "Nombre y apellidos", "Cargo", "Tipo de contrato", "Correo de acceso", "Permisos recomendados")
    departments = (
        ("Dirección", "Gestión general del centro y finanzas"),
        ("Finanzas / Contabilidad", "Finanzas y facturación solamente"),
        ("Administración / Secretaría", "Matrículas, alumnado y convocatorias; sin finanzas"),
        ("Formación / Coordinación", "Cursos, docentes, aulas y convocatorias; sin finanzas"),
        ("Comercial / Admisiones", "Leads y admisiones; sin finanzas"),
        ("Comunicación / Publicidad", "Web y publicidad; sin finanzas ni convocatorias"),
    )
    table = document.add_table(rows=1 + len(departments) * 2, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, (3.3, 4.7, 3.5, 3.2, 4.7, 6.4))
    set_repeat_header(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, BLUE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    row_index = 1
    for department, permissions in departments:
        for person_index in range(2):
            row = table.rows[row_index]
            row.cells[0].text = department if person_index == 0 else ""
            row.cells[1].text = " "
            row.cells[2].text = " "
            row.cells[3].text = " "
            row.cells[4].text = " "
            row.cells[5].text = permissions if person_index == 0 else ""
            if person_index == 0:
                set_cell_shading(row.cells[0], LIGHT_GRAY)
                set_cell_shading(row.cells[5], LIGHT_GRAY)
            row_index += 1

    style_table_text(table, header_rows=1, label_columns=(0,))

    note = document.add_paragraph()
    style_paragraph(note, before=6)
    run = note.add_run("Si una persona necesita permisos diferentes, indíquelo directamente en la última columna.")
    set_run(run, size=8.5, color=MUTED)
    document.save(output_path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for center_name, slug in CENTERS:
        build_company_form(center_name, OUTPUT_DIR / f"{slug}-01-datos-empresa-simplificado.docx")
        build_staff_form(center_name, OUTPUT_DIR / f"{slug}-02-personal-y-accesos-simplificado.docx")


if __name__ == "__main__":
    main()
